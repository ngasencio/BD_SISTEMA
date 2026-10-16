"""Informe por correo de la actualización automática (`services_informe_actualizacion`).

Cubre lo pedido el 2026-10-16:
- «En trámite hace más de 10 días» sale con los MÁS RECIENTES primero (antes: los más antiguos, que
  llegaban a formularios de 2022) y sin el recorte a 200 de `_diff_fsc`.
- El correo tiene diseño (cabecera con logo incrustado, tarjetas, etiquetas de estado) sin perder datos,
  escapa todo lo que viene de la base y no pasa del límite de tamaño de Gmail.
"""
import io
from datetime import date, datetime

from django.test import SimpleTestCase

from api import services_informe_actualizacion as inf
from api.services_actualizacion_diaria import ResultadoEtapa

HOY = date(2026, 10, 16)
CONTEXTO = {'inicio': datetime(2026, 10, 16, 7, 30), 'segundos': 2217,
            'ventana': 'OC y Licitaciones: últimos 5 días + hoy (11-10-2026 al 16-10-2026); '
                       'Compras Ágiles: ayer y hoy (15-10-2026 al 16-10-2026)'}


def _fsc(folio, dias, estado='P', anho=2026, unidad='DEPTO X'):
    return {'folio': folio, 'anho': anho, 'unidad_requirente': unidad, 'fecha_solicitud': f'2026-01-{folio % 28 + 1:02d}',
            'estado': estado, 'monto_estimado': 1000 * folio, 'usuario_requirente': 'PERSONA',
            'destino_actual': None, 'dias': dias}


def _clave(r):
    return (r['folio'], r['anho'], r['unidad_requirente'] or '', r['fecha_solicitud'] or '')


def _foto(filas):
    return {_clave(r): r for r in filas}


def _pegados(resultado):
    return next(l for l in resultado['listas'] if l['titulo'].startswith('En trámite'))


class FscEnTramiteMasRecientesPrimeroTest(SimpleTestCase):
    def _comparar(self, filas):
        foto = _foto(filas)
        return inf.comparar_fsc((dict(foto), set(), HOY, None), (foto, set(), HOY, None))

    def test_los_mas_recientes_van_primero_y_los_antiguos_al_final(self):
        r = self._comparar([_fsc(1, 1500), _fsc(2, 11), _fsc(3, 400), _fsc(4, 35)])
        lista = _pegados(r)
        self.assertEqual([f[-1] for f in lista['filas']], [11, 35, 400, 1500])   # columna «Días»
        self.assertIn('más recientes primero', lista['titulo'])
        self.assertNotIn('antiguos', lista['titulo'])

    def test_no_se_recorta_a_200_asi_los_recientes_nunca_quedan_fuera(self):
        # `_diff_fsc` los ordena de más antiguo a más reciente y se queda con 200: con 250 en trámite
        # perdería justo los 50 más recientes.
        r = self._comparar([_fsc(i, 11 + i) for i in range(250)])
        lista = _pegados(r)
        self.assertEqual(lista['total'], 250)
        self.assertEqual(len(lista['filas']), 250)
        self.assertEqual(lista['filas'][0][-1], 11)
        self.assertEqual(dict(r['resumen'])['En trámite hace más de 10 días'], 250)

    def test_umbral_es_estricto_y_excluye_derivados_rechazados_y_sin_fecha(self):
        r = self._comparar([_fsc(1, 10), _fsc(2, 11), _fsc(3, 500, estado='AC'), _fsc(4, 500, estado='R'),
                            _fsc(5, None)])
        self.assertEqual([f[0] for f in _pegados(r)['filas']], ['2/2026'])

    def test_a_igual_antiguedad_el_orden_es_estable(self):
        r = self._comparar([_fsc(5, 20), _fsc(9, 20), _fsc(7, 20, anho=2025)])
        self.assertEqual([f[0] for f in _pegados(r)['filas']], ['9/2026', '5/2026', '7/2025'])


class EtiquetasDeEstadoTest(SimpleTestCase):
    def test_en_fsc_solo_ac_cierra_el_tramite_y_dc_es_intermedio(self):
        fsc = inf._BANDEJA_FINAL['formularios']
        self.assertEqual(inf._tono_estado('AC · A Comprador', fsc), 'po')
        self.assertEqual(inf._tono_estado('DC · Derivación Compras', fsc), 'pi')
        self.assertEqual(inf._tono_estado('AA · Autorizador Abastecimiento', fsc), 'pi')

    def test_en_pac_dc_es_la_ficha_aprobada(self):
        self.assertEqual(inf._tono_estado('DC · Derivación Compras', inf._BANDEJA_FINAL['formularios_pac']), 'po')

    def test_rechazos_y_estados_de_mercado_publico(self):
        self.assertEqual(inf._tono_estado('R · Rechazado'), 'pb')
        self.assertEqual(inf._tono_estado('Desierta (o art. 3 ó 9 Ley 19.886)'), 'pb')
        self.assertEqual(inf._tono_estado('Cancelada'), 'pb')
        self.assertEqual(inf._tono_estado('Recepción Conforme'), 'po')
        self.assertEqual(inf._tono_estado('Adjudicada'), 'po')
        self.assertEqual(inf._tono_estado('Publicada'), 'pi')
        self.assertEqual(inf._tono_estado(''), 'pn')

    def test_dias_de_atraso_por_tramo(self):
        self.assertEqual(
            [inf._tono_dias(d) for d in (11, 30, 31, 90, 91, 1400, None)],
            ['pw', 'pw', 'pr', 'pr', 'pb', 'pb', 'pn'])


def _etapa(etapa, listas=(), **kw):
    base = dict(etapa=etapa, estado='OK', detalle='listo', segundos=65, rango='Últimos 5 días + hoy',
                nuevos=3, cambios=1, total=100, resumen=[('Registros en la base de datos', 100), ('Nuevos', 3)],
                listas=list(listas), notas=[])
    base.update(kw)
    return ResultadoEtapa(**base)


def _todas(**sobre):
    return [_etapa(e, **sobre.get(e, {})) for e in inf.ORDEN_INFORME]


class DisenoDelCorreoTest(SimpleTestCase):
    def test_documento_completo_con_logo_tarjetas_y_pie(self):
        html = inf.construir_informe(_todas(), CONTEXTO).html
        self.assertTrue(html.startswith('<!DOCTYPE html>'))
        self.assertIn('cid:logo', html)                       # logo incrustado, no de internet
        self.assertNotIn('http://', html)
        self.assertNotIn('https://', html)
        for texto in ('Informe de actualización de datos', 'Resumen por bloque', 'Detalle de cada fuente',
                      'Corrida del 16-10-2026 a las 07:30', 'Servicio de Salud Osorno', 'Gobierno de Chile',
                      'OC y Licitaciones: últimos 5 días + hoy (11-10-2026 al 16-10-2026)',
                      'Compras Ágiles: ayer y hoy (15-10-2026 al 16-10-2026)', 'Resultado: OK'):
            self.assertIn(texto, html)

    def test_tarjetas_resumen_suman_todos_los_bloques(self):
        html = inf.construir_informe(_todas(), CONTEXTO).html
        self.assertIn('5 de 5', html)
        self.assertIn('>15<', html)    # 5 bloques × 3 nuevos
        self.assertIn('>5<', html)     # 5 bloques × 1 cambio de estado

    def test_un_error_se_ve_en_la_cabecera_la_tarjeta_y_el_bloque(self):
        res = _todas(ordenes_compra=dict(estado='ERROR', detalle='API sin respuesta tras 3 intentos',
                                         nuevos=None, cambios=None, total=None, resumen=[]))
        html = inf.construir_informe(res, CONTEXTO).html
        self.assertIn('Resultado: CON ERRORES', html)
        self.assertIn('4 de 5', html)
        self.assertIn('API sin respuesta tras 3 intentos', html)
        self.assertIn('#fde8ea', html)

    def test_etapa_no_incluida_se_indica_sin_romper(self):
        res = [r for r in _todas() if r.etapa != 'compra_agil']
        html = inf.construir_informe(res, CONTEXTO).html
        self.assertIn('no incluida en esta corrida', html)
        self.assertIn('Compras Ágiles', html)

    def test_todo_lo_que_viene_de_la_base_se_escapa(self):
        lst = inf.lista('Lista', ['Código', 'Nombre', 'Estado'],
                        [['<b>X</b>', '<script>alert(1)</script>', '<img src=x onerror=alert(1)>']])
        res = _todas(ordenes_compra=dict(listas=[lst], detalle='<i>detalle</i>'))
        html = inf.construir_informe(res, CONTEXTO).html
        for peligroso in ('<script>', '<img src=x', '<b>X</b>', '<i>detalle</i>'):
            self.assertNotIn(peligroso, html)
        self.assertIn('&lt;script&gt;alert(1)&lt;/script&gt;', html)

    def test_entidades_html_que_ya_vienen_en_los_datos_no_se_muestran_literales(self):
        # Mercado Público entrega «BARRÍA &amp; MEDINA»: debe leerse «BARRÍA & MEDINA».
        self.assertEqual(inf._texto('BARRÍA &amp; MEDINA'), 'BARRÍA & MEDINA')
        lst = inf.lista('Lista', ['Código', 'Proveedor'], [['A-1', inf._texto('BARRÍA &amp; MEDINA')]])
        html = inf.construir_informe(_todas(ordenes_compra=dict(listas=[lst])), CONTEXTO).html
        self.assertIn('BARRÍA &amp; MEDINA', html)       # una sola vez escapado
        self.assertNotIn('&amp;amp;', html)

    def test_columnas_se_dibujan_segun_su_tipo(self):
        lst = inf.lista('FSC', ['Folio/Año', 'Unidad', 'Estado', 'Monto estimado', 'Días'],
                        [['1/2026', 'DEPTO', 'AC · A Comprador', '$ 1.000', 11]])
        html = inf.construir_informe(_todas(formularios=dict(listas=[lst])), CONTEXTO).html
        self.assertIn('<td class="c">1/2026</td>', html)                    # código en monoespaciada
        self.assertIn('<td class="n">$ 1.000</td>', html)                   # monto a la derecha
        self.assertIn('<span class="p po">AC · A Comprador</span>', html)   # estado como etiqueta
        self.assertIn('<span class="p pw">11</span>', html)                 # atraso por tramo

    def test_cuerpo_corto_mantiene_todas_las_filas_por_lista(self):
        lst = inf.lista('Nuevos', ['Código', 'Nombre'], [[f'C-{i}', 'x'] for i in range(25)])
        html = inf.construir_informe(_todas(ordenes_compra=dict(listas=[lst])), CONTEXTO).html
        self.assertIn('C-24', html)
        self.assertNotIn('el detalle completo está en el Excel adjunto', html)

    def test_si_el_html_pasa_el_limite_de_gmail_se_reduce_y_el_excel_conserva_todo(self):
        largo = 'Descripción larga de un producto o servicio de prueba ' * 3
        listas = [inf.lista(f'Lista {i}', ['Código', 'Nombre', 'Unidad', 'Proveedor', 'Estado'],
                            [[f'C-{i}-{k}', largo[:90], largo[:60], largo[:60], 'Recepción Conforme']
                             for k in range(300)]) for i in range(12)]
        informe = inf.construir_informe(_todas(ordenes_compra=dict(listas=listas)), CONTEXTO)
        self.assertLessEqual(len(informe.html.encode('utf-8')), inf.LIMITE_HTML_BYTES)
        self.assertIn('el detalle completo está en el Excel adjunto', informe.html)
        self.assertNotIn('C-0-299', informe.html)
        # El Excel adjunto trae las 300 filas (+ encabezado) de la primera lista.
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(informe.adjuntos[0][1]))
        hoja = next(wb[n] for n in wb.sheetnames if n.endswith('Lista 0'))
        self.assertEqual(hoja.max_row, 301)

    def test_aviso_de_omitida_con_el_mismo_diseno(self):
        informe = inf.informe_omitido(CONTEXTO, 'ya había otra actualización <en curso>')
        self.assertIn('OMITIDA', informe.asunto)
        self.assertIn('cid:logo', informe.html)
        self.assertIn('Ya había otra actualización &lt;en curso&gt;', informe.html)
        self.assertNotIn('En el asunto', informe.html)


class MensajeConLogoIncrustadoTest(SimpleTestCase):
    def _crudo(self):
        informe = inf.construir_informe(_todas(), CONTEXTO)
        msg = inf._clase_mensaje()(informe.asunto, informe.texto, None, ['informe@test.local'])
        msg.attach_alternative(informe.html, 'text/html')
        for nombre, datos, mime in informe.adjuntos:
            msg.attach(nombre, datos, mime)
        return msg.message().as_string()

    def test_el_logo_viaja_junto_al_html_y_el_excel_va_aparte(self):
        crudo = self._crudo()
        self.assertIn('Content-ID: <logo>', crudo)
        self.assertIn('multipart/related', crudo)       # logo + html juntos (Outlook web lo exige)
        self.assertIn('multipart/alternative', crudo)
        self.assertIn('multipart/mixed', crudo)         # el Excel adjunto
        self.assertIn('.xlsx', crudo)

    def test_si_el_logo_no_se_puede_leer_el_informe_sale_igual(self):
        from unittest import mock
        with mock.patch('api.services_notificacion_plan._bytes_logo', side_effect=OSError('sin logo')):
            self.assertIs(inf._clase_mensaje(), inf.EmailMultiAlternatives)
