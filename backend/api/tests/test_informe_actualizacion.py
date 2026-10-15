"""Informe por correo de cada actualización automática: diferencias antes/después por bloque
(OC, Licitaciones, Compras Ágiles, Panel FSC, Panel PAC) y construcción del correo.

Las diferencias se prueban con filas reales de la base de pruebas; el correo con resultados
armados a mano. No se envía nada: solo se construye (y `enviar_informe` va a la bandeja en
memoria de settings_test)."""
import io
from datetime import date, datetime

from django.core import mail
from django.test import SimpleTestCase, TestCase, override_settings

from api import services_informe_actualizacion as inf
from api.models import (
    CompraAgilResumen, FormularioFSC, FormularioFSCDerivado, FormularioPAC, Licitacion, OrdenCompra,
)
from api.services_actualizacion_diaria import ResultadoEtapa


def _por_titulo(cmp_):
    return {lst['titulo']: lst for lst in cmp_['listas']}


class DiferenciasPorBloqueTest(TestCase):
    def test_oc_nuevas_y_cambios_de_estado(self):
        OrdenCompra.objects.create(codigo_oc='1-1-SE26', EstadoOC='Enviada a proveedor', NombreOC='Insumos',
                                   P_Nombre='Proveedor A', C_Unidad='Bodega', TotalNeto=100000)
        OrdenCompra.objects.create(codigo_oc='2-2-SE26', EstadoOC='Aceptada', NombreOC='Otra')
        antes = inf.tomar_oc()

        OrdenCompra.objects.filter(codigo_oc='1-1-SE26').update(EstadoOC='Aceptada')
        OrdenCompra.objects.create(codigo_oc='3-3-SE26', EstadoOC='Enviada a proveedor', NombreOC='Nueva',
                                   P_Nombre='Proveedor B', TotalNeto=250000,
                                   FechaEnvio=datetime(2026, 10, 14, 9, 0))
        c = inf.comparar_oc(antes, inf.tomar_oc())

        self.assertEqual((c['nuevos'], c['cambios'], c['total']), (1, 1, 3))
        listas = _por_titulo(c)
        self.assertEqual(listas['OC nuevas']['filas'][0][0], '3-3-SE26')
        self.assertEqual(listas['OC nuevas']['filas'][0][5], '$ 250.000')
        fila = listas['OC con cambio de estado']['filas'][0]
        self.assertEqual((fila[0], fila[-2], fila[-1]), ('1-1-SE26', 'Enviada a proveedor', 'Aceptada'))
        self.assertIn(('OC nuevas', 1), c['resumen'])

    def test_oc_sin_cambios_da_ceros_y_listas_vacias(self):
        OrdenCompra.objects.create(codigo_oc='1-1-SE26', EstadoOC='Aceptada')
        antes = inf.tomar_oc()
        c = inf.comparar_oc(antes, inf.tomar_oc())
        self.assertEqual((c['nuevos'], c['cambios']), (0, 0))
        self.assertTrue(all(lst['total'] == 0 for lst in c['listas']))

    def test_licitaciones_nuevas_cambios_y_adjudicadas(self):
        Licitacion.objects.create(codigo_licitacion='L-1', Estado='Publicada', Nombre='Una', Numero='1')
        Licitacion.objects.create(codigo_licitacion='L-2', Estado='Publicada', Nombre='Dos', Numero='2')
        antes = inf.tomar_licitaciones()
        Licitacion.objects.filter(codigo_licitacion='L-1').update(Estado='Adjudicada')
        Licitacion.objects.filter(codigo_licitacion='L-2').update(Estado='Cerrada')
        Licitacion.objects.create(codigo_licitacion='L-3', Estado='Publicada', Nombre='Tres', Numero='3')
        c = inf.comparar_licitaciones(antes, inf.tomar_licitaciones())
        self.assertEqual((c['nuevos'], c['cambios'], c['total']), (1, 2, 3))
        self.assertIn(('De ellas, adjudicadas', 1), c['resumen'])

    def test_compra_agil_nuevas_cambios_y_oc_vinculada(self):
        CompraAgilResumen.objects.create(codigocompraagil='1057532-1-COT26', estadoglosa='Publicada',
                                         nombre='Insumos', unidadcompra='Bodega')
        CompraAgilResumen.objects.create(codigocompraagil='1057532-2-COT26', estadoglosa='Proveedor seleccionado',
                                         nombre='Otros', oc_codigo='')
        antes = inf.tomar_compra_agil()
        CompraAgilResumen.objects.filter(codigocompraagil='1057532-1-COT26').update(estadoglosa='Cerrada')
        CompraAgilResumen.objects.filter(codigocompraagil='1057532-2-COT26').update(oc_codigo='1057532-9-AG26')
        CompraAgilResumen.objects.create(codigocompraagil='1057532-3-COT26', estadoglosa='Publicada', nombre='Nueva')
        c = inf.comparar_compra_agil(antes, inf.tomar_compra_agil())
        self.assertEqual((c['nuevos'], c['cambios'], c['total']), (1, 1, 3))
        self.assertIn(('Con OC recién vinculada', 1), c['resumen'])
        self.assertEqual(_por_titulo(c)['Compras ágiles con OC recién vinculada']['filas'][0][0],
                         '1057532-2-COT26')

    def test_pac_nuevas_cambios_y_eliminadas(self):
        FormularioPAC.objects.create(folio=1, anho=2026, estado='P', departamento='Depto A', monto=1000)
        FormularioPAC.objects.create(folio=2, anho=2026, estado='FR', departamento='Depto B')
        antes = inf.tomar_pac()
        FormularioPAC.objects.filter(folio=1).update(estado='FR')
        FormularioPAC.objects.filter(folio=2).delete()
        FormularioPAC.objects.create(folio=3, anho=2026, estado='P', departamento='Depto C', monto=-500)
        c = inf.comparar_pac(antes, inf.tomar_pac())
        self.assertEqual((c['nuevos'], c['cambios'], c['total']), (1, 1, 2))
        listas = _por_titulo(c)
        self.assertEqual(listas['Fichas PAC que cambiaron de bandeja']['filas'][0][-1], 'P · Pendiente Firmas')
        self.assertEqual(listas['Fichas PAC que ya no están en el Panel']['filas'][0][0], '2/2026')

    def test_fsc_nuevos_cambios_y_derivados(self):
        FormularioFSC.objects.create(folio=10, anho=2026, unidad_requirente='UNIDAD A', estado='P',
                                     fecha_solicitud='2026-10-01', monto_estimado=5000)
        antes = inf.tomar_fsc()
        FormularioFSC.objects.filter(folio=10).update(estado='DC')
        FormularioFSC.objects.create(folio=11, anho=2026, unidad_requirente='UNIDAD B', estado='P',
                                     fecha_solicitud='2026-10-14', monto_estimado=1)
        FormularioFSCDerivado.objects.create(folio=10, anho=2026, unidad_requirente='UNIDAD A',
                                             fecha_solicitud='2026-10-01')
        c = inf.comparar_fsc(antes, inf.tomar_fsc())
        self.assertEqual((c['nuevos'], c['cambios'], c['total']), (1, 1, 2))
        listas = _por_titulo(c)
        self.assertEqual(listas['FSC que cambiaron de bandeja']['filas'][0][-1], 'P · Pendiente Firmas')
        self.assertEqual(listas['Derivados a comprador (nuevos)']['total'], 1)


def _resultado(etapa, estado='OK', nuevos=0, cambios=0, total=0, **kw):
    r = ResultadoEtapa(etapa, estado, kw.pop('detalle', 'detalle'), kw.pop('segundos', 65.0))
    r.nuevos, r.cambios, r.total = nuevos, cambios, total
    for k, v in kw.items():
        setattr(r, k, v)
    return r


class ConstruccionDelCorreoTest(SimpleTestCase):
    CTX = {'inicio': datetime(2026, 10, 15, 7, 30), 'segundos': 1500, 'ventana': 'Últimos 5 días + hoy'}

    def _filas(self, n):
        return [[f'C-{i}', f'Nombre {i}', 'Unidad'] for i in range(n)]

    def test_bloques_en_el_orden_pedido_y_los_no_incluidos_se_marcan(self):
        res = [
            _resultado('licitaciones', nuevos=2, cambios=1, total=834),
            _resultado('ordenes_compra', nuevos=12, cambios=8, total=19500),
            _resultado('formularios', nuevos=1, total=13615),
        ]
        informe = inf.construir_informe(res, self.CTX)
        html = informe.html
        # Orden del correo: OC, Panel FSC, Panel PAC, Compras Ágiles, Licitaciones
        posiciones = [html.index(t) for t in ('Órdenes de Compra', 'Panel FSC', 'Panel PAC',
                                              'Compras Ágiles', 'Licitaciones')]
        self.assertEqual(posiciones, sorted(posiciones))
        self.assertIn('no se incluye en esta corrida', html)  # PAC y Compras Ágiles
        self.assertIn('OC +12/8', informe.asunto)
        self.assertIn('LIC +2/1', informe.asunto)
        self.assertIn('07:30', informe.asunto)
        self.assertIn('OK', informe.asunto)

    def test_un_error_se_ve_en_asunto_cuerpo_y_texto(self):
        res = [_resultado('ordenes_compra', 'ERROR', detalle='La API no respondió'),
               _resultado('licitaciones', nuevos=1, total=10)]
        informe = inf.construir_informe(res, self.CTX)
        self.assertIn('CON ERRORES', informe.asunto)
        self.assertIn('OC ERROR', informe.asunto)
        self.assertIn('CON ERROR', informe.html)
        self.assertIn('La API no respondió', informe.texto)

    def test_las_listas_largas_se_recortan_y_el_excel_trae_todo(self):
        from openpyxl import load_workbook
        lst = inf.lista('OC nuevas', ['Código', 'Nombre', 'Unidad'], self._filas(60))
        res = [_resultado('ordenes_compra', nuevos=60, total=60, listas=[lst],
                          resumen=[('OC nuevas', 60)], notas=['Nota de prueba'])]
        informe = inf.construir_informe(res, self.CTX)
        self.assertIn('y 35 más', informe.html)  # 60 - 25 mostradas
        self.assertIn('C-24', informe.html)
        self.assertNotIn('C-25', informe.html)
        self.assertIn('Nota de prueba', informe.html)

        nombre, datos, mime = informe.adjuntos[0]
        self.assertTrue(nombre.startswith('actualizacion_20261015_0730'))
        self.assertIn('spreadsheetml', mime)
        libro = load_workbook(io.BytesIO(datos))
        self.assertEqual(libro.sheetnames[0], 'Resumen')
        hoja = libro[[n for n in libro.sheetnames if 'OC nuevas' in n][0]]
        self.assertEqual(hoja.max_row, 61)  # encabezado + 60 filas

    def test_el_html_escapa_el_contenido(self):
        lst = inf.lista('OC nuevas', ['Código', 'Nombre'], [['X', '<script>alert(1)</script>']])
        res = [_resultado('ordenes_compra', nuevos=1, total=1, listas=[lst], detalle='<b>detalle</b>')]
        html = inf.construir_informe(res, self.CTX).html
        self.assertNotIn('<script>', html)
        self.assertIn('&lt;script&gt;', html)
        self.assertNotIn('<b>detalle</b>', html)

    @override_settings(SYNC_INFORME_A=['nicolas.asencio@redsalud.gob.cl'])
    def test_se_envia_a_los_destinatarios_configurados_con_adjunto(self):
        informe = inf.construir_informe([_resultado('licitaciones', nuevos=1, total=5)], self.CTX)
        destinos = inf.enviar_informe(informe)
        self.assertEqual(destinos, ['nicolas.asencio@redsalud.gob.cl'])
        correo = mail.outbox[0]
        self.assertEqual(correo.to, ['nicolas.asencio@redsalud.gob.cl'])
        self.assertEqual(len(correo.attachments), 1)
        self.assertEqual(correo.alternatives[0][1], 'text/html')

    @override_settings(SYNC_INFORME_A=[])
    def test_sin_destinatarios_lanza(self):
        informe = inf.construir_informe([_resultado('licitaciones')], self.CTX)
        with self.assertRaises(RuntimeError):
            inf.enviar_informe(informe)

    def test_informe_omitido_es_corto(self):
        o = inf.informe_omitido(self.CTX, 'Sigue habiendo una actualización en curso.')
        self.assertIn('OMITIDA', o.asunto)
        self.assertEqual(o.adjuntos, [])

    def test_el_destinatario_por_defecto_es_el_pedido(self):
        from django.conf import settings
        self.assertIn('nicolas.asencio@redsalud.gob.cl', settings.SYNC_INFORME_A)
